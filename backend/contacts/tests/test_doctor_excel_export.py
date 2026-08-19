"""Excel export tests for the Doctor endpoint.

Covers: /api/v1/contacts/doctors/?export=excel — the export the دفترچه تماس
(contacts directory) page's "خروجی اکسل" button links to, since that page
is entirely JS/API-driven (contacts_directory.js) rather than a Django
admin ChangeList.
"""

import io

from django.contrib.auth.models import Group, User
from rest_framework import status
from rest_framework.test import APITestCase

from contacts.models import CooperationStatus, Doctor, DoctorSpecialty

URL     = '/api/v1/contacts/doctors/'
XLSX_CT = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _admin():
    u = User.objects.create_user(username=f'doc_adm_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='admin')
    u.groups.add(g)
    return u


def _specialty(name=None):
    import random
    return DoctorSpecialty.objects.create(name=name or f'تخصص_{random.randint(1000, 9999)}')


def _doctor(name, specialty, cooperation_status=CooperationStatus.ACTIVE, is_active=True):
    import random
    return Doctor.objects.create(
        full_name=name,
        specialty=specialty,
        phone_number=f'0910{random.randint(1000000, 9999999)}',
        cooperation_status=cooperation_status,
        is_active=is_active,
    )


def _open_workbook(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


def _header_row_index(ws, *, contains):
    for row in range(1, ws.max_row + 1):
        if any(ws.cell(row, c).value == contains for c in range(1, ws.max_column + 1)):
            return row
    raise AssertionError(f'header row containing {contains!r} not found')


def _header_labels(ws, header_row):
    return [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]


def _column_values(ws, header_row, col_idx):
    return [
        ws.cell(row, col_idx).value
        for row in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row, col_idx).value not in (None, '')
    ]


class DoctorExcelExportTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        self.spec1 = _specialty('جراحی عمومی')
        self.spec2 = _specialty('ارتوپدی')
        self.d1 = _doctor('دکتر رضا احمدی', self.spec1, cooperation_status=CooperationStatus.ACTIVE)
        self.d2 = _doctor('دکتر سارا محمدی', self.spec2, cooperation_status=CooperationStatus.INACTIVE)

    def _get(self, **params):
        return self.client.get(URL, {'export': 'excel', **params})

    def _name_column_values(self, ws):
        header_row = _header_row_index(ws, contains='نام و نام خانوادگی')
        name_col   = _header_labels(ws, header_row).index('نام و نام خانوادگی') + 1
        return _column_values(ws, header_row, name_col)

    def test_returns_xlsx_content_type(self):
        r = self._get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_content_disposition_has_filename(self):
        r = self._get()
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('.xlsx', r['Content-Disposition'])

    def test_rtl_worksheet(self):
        r = self._get()
        ws = _open_workbook(r.content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_persian_headers_present(self):
        r = self._get()
        ws = _open_workbook(r.content).active
        header_row = _header_row_index(ws, contains='نام و نام خانوادگی')
        headers = _header_labels(ws, header_row)
        self.assertIn('نام و نام خانوادگی', headers)
        self.assertIn('تخصص',               headers)
        self.assertIn('وضعیت همکاری',       headers)
        # Financial fields (commission %, rate per surgery) must never
        # appear in this general directory export.
        self.assertNotIn('درصد کمیسیون مرکز', headers)
        self.assertNotIn('نرخ هر عمل',        headers)

    def test_no_financial_fields_in_any_cell(self):
        # Belt-and-suspenders: scan every cell, not just the header row.
        r = self._get()
        ws = _open_workbook(r.content).active
        all_vals = [
            ws.cell(row, col).value
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertNotIn('درصد کمیسیون مرکز', all_vals)
        self.assertNotIn('نرخ هر عمل',        all_vals)

    def test_cooperation_status_filter_applied(self):
        r = self._get(cooperation_status='active')
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('دکتر رضا احمدی', names)
        self.assertNotIn('دکتر سارا محمدی', names)

    def test_specialty_filter_applied(self):
        r = self._get(specialty=self.spec1.id)
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('دکتر رضا احمدی', names)
        self.assertNotIn('دکتر سارا محمدی', names)

    def test_search_filter_applied(self):
        r = self._get(search='رضا')
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('دکتر رضا احمدی', names)
        self.assertNotIn('دکتر سارا محمدی', names)

    def test_empty_result_is_valid_workbook(self):
        r = self._get(search='نام‌ نامرتبط با هیچ‌کس')
        self.assertEqual(r.status_code, 200)
        wb = _open_workbook(r.content)
        self.assertIsNotNone(wb)

    def test_unauthenticated_returns_401(self):
        self.client.force_authenticate(None)
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_401_UNAUTHORIZED)
