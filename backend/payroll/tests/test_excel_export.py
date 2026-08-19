"""Excel export tests for payroll endpoints (CLI-54).

Covers: /api/v2/payroll/reports/employee-cost/?export=excel
"""

import io
import datetime
from decimal import Decimal

from django.contrib.auth.models import Group, User
from rest_framework import status
from rest_framework.test import APITestCase

from employees.models import Employee, JobPosition
from payroll.models import MonthlyWage

URL     = '/api/v2/payroll/reports/employee-cost/'
XLSX_CT = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _admin():
    u = User.objects.create_user(username=f'pay_adm_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='admin')
    u.groups.add(g)
    return u


def _finance():
    u = User.objects.create_user(username=f'pay_fin_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='finance_user')
    u.groups.add(g)
    return u


def _position():
    import random
    return JobPosition.objects.create(name=f'موقعیت_{random.randint(10000, 99999)}')


def _employee(name, position):
    import random
    return Employee.objects.create(
        full_name=name,
        national_id=str(random.randint(1000000000, 9999999999)),
        gender='female',
        job_position=position,
        start_date=datetime.date(2021, 1, 1),
        personal_phone='09100000000',
        emergency_contact_phone='09200000000',
    )


def _wage(employee, amount, start=datetime.date(2025, 1, 1), end=None):
    return MonthlyWage.objects.create(
        employee=employee, amount=amount, start_date=start, end_date=end,
    )


def _open_workbook(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class EmployeeCostReportExcelTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        self.pos = _position()
        self.emp = _employee('زهرا مرادی', self.pos)
        _wage(self.emp, Decimal('3000000'))

    def _get(self, **params):
        return self.client.get(URL, {'export': 'excel', **params})

    def test_returns_xlsx_content_type(self):
        r = self._get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_content_disposition_filename(self):
        r = self._get()
        self.assertIn('.xlsx', r['Content-Disposition'])

    def test_rtl_worksheet(self):
        r = self._get()
        ws = _open_workbook(r.content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_persian_headers(self):
        r = self._get()
        ws = _open_workbook(r.content).active
        # Meta rows precede the header; collect ALL cell values and check
        all_vals = [ws.cell(row, col).value for row in range(1, ws.max_row + 1)
                    for col in range(1, ws.max_column + 1) if ws.cell(row, col).value]
        self.assertIn('کارمند',        all_vals)
        self.assertIn('حقوق ثابت',     all_vals)
        self.assertIn('مجموع پرداختی', all_vals)

    def test_employee_data_in_rows(self):
        r = self._get(start_date='2025-01-01', end_date='2025-12-31')
        ws = _open_workbook(r.content).active
        all_vals = [ws.cell(row, col).value for row in range(1, ws.max_row + 1)
                    for col in range(1, ws.max_column + 1) if ws.cell(row, col).value]
        self.assertIn('زهرا مرادی', all_vals)

    def test_no_pagination_all_rows_exported(self):
        pos = _position()
        for i in range(5):
            emp = _employee(f'کارمند_اضافه_{i}', pos)
            _wage(emp, Decimal('1000000'))
        r = self._get(page_size='2')
        ws = _open_workbook(r.content).active
        # All 6 employees (1 setUp + 5 new) should be present
        # Find header row first
        header_row = None
        for row in range(1, ws.max_row + 1):
            if ws.cell(row, 1).value == 'کارمند':
                header_row = row
                break
        self.assertIsNotNone(header_row)
        data_rows = [ws.cell(row, 1).value for row in range(header_row + 1, ws.max_row + 1)
                     if ws.cell(row, 1).value]
        self.assertGreaterEqual(len(data_rows), 6)

    def test_finance_user_can_export(self):
        self.client.force_authenticate(_finance())
        r = self._get()
        self.assertEqual(r.status_code, 200)

    def test_plain_user_gets_403(self):
        u = User.objects.create_user(username='plain2', password='x')
        self.client.force_authenticate(u)
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    def test_wage_type_filter_applied(self):
        r = self._get(wage_type='monthly')
        self.assertEqual(r.status_code, 200)
        # Response should be Excel, not error
        self.assertEqual(r['Content-Type'], XLSX_CT)
