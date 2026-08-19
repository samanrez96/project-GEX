"""Excel export tests for the employees endpoint (CLI-54)."""

import io
import datetime
from decimal import Decimal

from django.contrib.auth.models import Group, User
from rest_framework import status
from rest_framework.test import APITestCase

from employees.models import Employee, JobPosition

URL   = '/api/v2/employees/'
XLSX_CT = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _admin():
    u = User.objects.create_user(username=f'emp_adm_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='admin')
    u.groups.add(g)
    return u


def _position(name=None):
    import random
    uid = random.randint(100000, 999999)
    return JobPosition.objects.create(name=name or f'Pos_{uid}')


def _employee(name, position, is_active=True, national_id=None):
    import random
    return Employee.objects.create(
        full_name=name,
        national_id=national_id or str(random.randint(1000000000, 9999999999)),
        gender='male',
        job_position=position,
        start_date=datetime.date(2022, 1, 1),
        personal_phone='09000000000',
        emergency_contact_phone='09111111111',
        is_active=is_active,
    )


def _open_workbook(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


def _header_row_index(ws, *, contains):
    for row in range(1, ws.max_row + 1):
        if any(ws.cell(row, c).value == contains for c in range(1, ws.max_column + 1)):
            return row
    raise AssertionError(f'header row containing {contains!r} not found')


def _header_labels(ws, header_row):
    return [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]


def _column_values(ws, header_row, col_idx):
    return [
        ws.cell(row, col_idx).value
        for row in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row, col_idx).value not in (None, '')
    ]


class EmployeeExcelExportTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        self.pos  = _position()
        self.pos2 = _position()
        self.e1 = _employee('علی رضایی',   self.pos,  is_active=True)
        self.e2 = _employee('مریم احمدی',  self.pos2, is_active=False)

    def _get(self, **params):
        return self.client.get(URL, {'export': 'excel', **params})

    def test_returns_xlsx_content_type(self):
        r = self._get(is_active='true')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_content_disposition_contains_filename(self):
        r = self._get(is_active='true')
        self.assertIn('attachment',    r['Content-Disposition'])
        self.assertIn('.xlsx',         r['Content-Disposition'])

    def test_workbook_openable(self):
        r = self._get(is_active='true')
        wb = _open_workbook(r.content)
        self.assertIsNotNone(wb)

    def test_rtl_worksheet(self):
        r = self._get(is_active='true')
        ws = _open_workbook(r.content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_persian_headers(self):
        r = self._get(is_active='true')
        ws = _open_workbook(r.content).active
        header_row = _header_row_index(ws, contains='نام کامل')
        headers = _header_labels(ws, header_row)
        self.assertIn('نام کامل',     headers)
        self.assertIn('پوزیشن شغلی',  headers)
        self.assertIn('شماره ملی',    headers)
        self.assertIn('وضعیت',        headers)

    def _name_column_values(self, ws):
        header_row = _header_row_index(ws, contains='نام کامل')
        name_col   = _header_labels(ws, header_row).index('نام کامل') + 1
        return _column_values(ws, header_row, name_col)

    def test_is_active_filter_applied(self):
        r = self._get(is_active='true')
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('علی رضایی',  names)
        self.assertNotIn('مریم احمدی', names)

    def test_all_employees_when_no_filter(self):
        r = self._get(is_active='all')
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('علی رضایی',  names)
        self.assertIn('مریم احمدی', names)

    def test_job_position_filter_applied(self):
        r = self._get(job_position=self.pos.id, is_active='all')
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('علی رضایی',    names)
        self.assertNotIn('مریم احمدی', names)

    def test_no_pagination_all_rows_exported(self):
        pos = _position()
        for i in range(5):
            _employee(f'کارمند_{i}', pos)
        r = self._get(is_active='all', page_size='2')
        ws = _open_workbook(r.content).active
        data_rows = self._name_column_values(ws)
        self.assertGreaterEqual(len(data_rows), 7)   # 2 setUp + 5 new

    def test_unauthenticated_returns_401(self):
        self.client.force_authenticate(None)
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_401_UNAUTHORIZED)
