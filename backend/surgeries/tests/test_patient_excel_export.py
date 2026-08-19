"""Excel export tests for the Patient admin changelist.

Covers: /admin/surgeries/patient/export-excel/ — the only export in this
project built on AdminExcelExportMixin (reusing the ModelAdmin's own
ChangeList), since the Patient list page is a genuine Django admin
changelist (unlike Surgery/Product/Purchase/Employee/Doctor, which are all
JS/API-driven pages exported through their DRF ViewSets instead).
"""

import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase

from surgeries.models import Patient

User = get_user_model()

EXPORT_URL = '/admin/surgeries/patient/export-excel/'


def _patient(full_name, case_code, internal_code=None):
    return Patient.objects.create(
        full_name=full_name,
        case_code=case_code,
        internal_code=internal_code,
        phone_number='09120000000',
        national_id='5555555555',
    )


def _open_workbook(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


def _header_row_index(ws, *, contains):
    for row in range(1, ws.max_row + 1):
        if any(ws.cell(row, c).value == contains for c in range(1, ws.max_column + 1)):
            return row
    raise AssertionError(f'header row containing {contains!r} not found')


def _header_labels(ws, header_row):
    return [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]


def _column_values(ws, header_row, col_idx):
    return [
        ws.cell(row, col_idx).value
        for row in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row, col_idx).value not in (None, '')
    ]


class PatientAdminExcelExportTest(TestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('pat_exp_admin', 'p@t.com', 'pass123')
        self.client.force_login(self.superuser)
        self.p1 = _patient('رضا کریمی', 'PAT001', internal_code='INT001')
        self.p2 = _patient('سارا احمدی', 'PAT002', internal_code='INT002')

    def _name_column_values(self, ws):
        header_row = _header_row_index(ws, contains='نام کامل')
        name_col   = _header_labels(ws, header_row).index('نام کامل') + 1
        return _column_values(ws, header_row, name_col)

    def test_returns_xlsx_content_type(self):
        r = self.client.get(EXPORT_URL)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def test_content_disposition_has_filename(self):
        r = self.client.get(EXPORT_URL)
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('.xlsx', r['Content-Disposition'])

    def test_rtl_worksheet(self):
        r = self.client.get(EXPORT_URL)
        ws = _open_workbook(r.content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_persian_headers_present(self):
        r = self.client.get(EXPORT_URL)
        ws = _open_workbook(r.content).active
        header_row = _header_row_index(ws, contains='نام کامل')
        headers = _header_labels(ws, header_row)
        self.assertIn('کد پرونده',       headers)
        self.assertIn('کد داخلی بیمار', headers)
        self.assertIn('نام کامل',        headers)
        self.assertIn('کد ملی',          headers)

    def test_search_query_reflected(self):
        r = self.client.get(EXPORT_URL, {'q': 'رضا'})
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('رضا کریمی',   names)
        self.assertNotIn('سارا احمدی', names)

    def test_no_search_returns_all(self):
        r = self.client.get(EXPORT_URL)
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('رضا کریمی',   names)
        self.assertIn('سارا احمدی', names)

    def test_empty_result_is_valid_workbook_not_500(self):
        r = self.client.get(EXPORT_URL, {'q': 'نام کاملا نامرتبط'})
        self.assertEqual(r.status_code, 200)
        wb = _open_workbook(r.content)
        self.assertIsNotNone(wb)

    def test_staff_without_view_permission_forbidden(self):
        staff = User.objects.create_user(
            'pat_exp_staff', 'staff@t.com', 'pass123', is_staff=True,
        )
        self.client.force_login(staff)
        r = self.client.get(EXPORT_URL)
        self.assertEqual(r.status_code, 403)

    def test_staff_with_view_permission_allowed(self):
        staff = User.objects.create_user(
            'pat_exp_staff2', 'staff2@t.com', 'pass123', is_staff=True,
        )
        perm = Permission.objects.get(codename='view_patient')
        staff.user_permissions.add(perm)
        self.client.force_login(staff)
        r = self.client.get(EXPORT_URL)
        self.assertEqual(r.status_code, 200)

    def test_unauthenticated_redirects_to_login(self):
        self.client.logout()
        r = self.client.get(EXPORT_URL)
        self.assertEqual(r.status_code, 302)
