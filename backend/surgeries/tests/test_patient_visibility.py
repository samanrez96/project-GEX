"""Tests for the hidden-Patient visibility feature.

Covers: Patient.objects.visible_to() / SurgeryHistory.objects.visible_to()
(the canonical, centralized visibility rule), PatientAdmin, PatientViewSet,
SurgeryHistoryViewSet, and forged-write protection.

"Main administrator" == Django's own is_superuser flag (accounts.permissions
.is_main_administrator) — there is no more explicit root-admin role in this
project (the "admin" group is a broad business role, not equivalent).
"""

import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from rest_framework.test import APITestCase

from surgeries.models import Patient, SurgeryHistory, SurgeryType

User = get_user_model()

PATIENTS_URL = '/api/v1/surgeries/patients/'
HISTORY_URL  = '/api/v1/surgeries/history/'


def _patient(case_code='PAT001', full_name='بیمار تست', hidden=False):
    p = Patient.objects.create(
        full_name=full_name, case_code=case_code,
        phone_number='09120000000', national_id='5555555555',
    )
    if hidden:
        p.is_hidden = True
        p.save(update_fields=['is_hidden'])
    return p


def _surgery_type():
    n = SurgeryType.objects.count()
    return SurgeryType.objects.create(
        name=f'عمومی_{n}', code=f'gen_{n}', base_rate=Decimal('1000000'),
    )


def _surgery(patient, surgery_type=None):
    return SurgeryHistory.objects.create(
        patient=patient,
        surgery_type=surgery_type or _surgery_type(),
        amount=Decimal('10000000'),
        surgery_date=datetime.datetime(2025, 6, 1, 10, 0, tzinfo=datetime.timezone.utc),
    )


def _staff_with_patient_perms():
    staff = User.objects.create_user(
        f'staff_{User.objects.count()}', 'staff@t.com', 'pass123', is_staff=True,
    )
    for codename in ('view_patient', 'change_patient', 'view_surgeryhistory'):
        staff.user_permissions.add(Permission.objects.get(codename=codename))
    return staff


# ---------------------------------------------------------------------------
# QuerySet / manager
# ---------------------------------------------------------------------------

class PatientVisibleToQuerySetTest(TestCase):

    def setUp(self):
        self.visible = _patient('VIS01', 'قابل مشاهده')
        self.hidden  = _patient('HID01', 'مخفی شده', hidden=True)
        self.superuser = User.objects.create_superuser('su1', 's@t.com', 'pass123')
        self.staff     = _staff_with_patient_perms()

    def test_superuser_sees_all(self):
        ids = set(Patient.objects.visible_to(self.superuser).values_list('pk', flat=True))
        self.assertEqual(ids, {self.visible.pk, self.hidden.pk})

    def test_normal_user_excludes_hidden(self):
        ids = set(Patient.objects.visible_to(self.staff).values_list('pk', flat=True))
        self.assertEqual(ids, {self.visible.pk})

    def test_anonymous_excludes_hidden(self):
        from django.contrib.auth.models import AnonymousUser
        ids = set(Patient.objects.visible_to(AnonymousUser()).values_list('pk', flat=True))
        self.assertEqual(ids, {self.visible.pk})

    def test_new_patient_defaults_to_visible(self):
        p = _patient('DEFAULT01', 'پیش‌فرض')
        self.assertFalse(p.is_hidden)

    def test_hide_sets_audit_fields(self):
        p = _patient('AUDIT01', 'حسابرسی')
        p.hide(self.superuser)
        p.refresh_from_db()
        self.assertTrue(p.is_hidden)
        self.assertEqual(p.hidden_by_id, self.superuser.pk)
        self.assertIsNotNone(p.hidden_at)

    def test_unhide_clears_audit_fields(self):
        p = _patient('AUDIT02', 'حسابرسی۲', hidden=True)
        p.hidden_by = self.superuser
        p.hidden_at = datetime.datetime.now(datetime.timezone.utc)
        p.save(update_fields=['hidden_by', 'hidden_at'])
        p.unhide()
        p.refresh_from_db()
        self.assertFalse(p.is_hidden)
        self.assertIsNone(p.hidden_by_id)
        self.assertIsNone(p.hidden_at)

    def test_hide_is_idempotent(self):
        p = _patient('IDEMP01', 'یکتوان', hidden=True)
        p.hide(self.superuser)   # already hidden — must not error
        p.refresh_from_db()
        self.assertTrue(p.is_hidden)

    def test_unhide_is_idempotent(self):
        p = _patient('IDEMP02', 'یکتوان۲')
        p.unhide()   # already visible — must not error
        p.refresh_from_db()
        self.assertFalse(p.is_hidden)


class SurgeryHistoryVisibleToQuerySetTest(TestCase):

    def setUp(self):
        self.visible_patient = _patient('SVIS01', 'بیمار مشهود')
        self.hidden_patient  = _patient('SHID01', 'بیمار مخفی', hidden=True)
        self.s_visible = _surgery(self.visible_patient)
        self.s_hidden  = _surgery(self.hidden_patient)
        self.superuser = User.objects.create_superuser('su2', 's2@t.com', 'pass123')
        self.staff     = _staff_with_patient_perms()

    def test_superuser_sees_all_surgeries(self):
        ids = set(SurgeryHistory.objects.visible_to(self.superuser).values_list('pk', flat=True))
        self.assertEqual(ids, {self.s_visible.pk, self.s_hidden.pk})

    def test_normal_user_excludes_hidden_patient_surgery(self):
        ids = set(SurgeryHistory.objects.visible_to(self.staff).values_list('pk', flat=True))
        self.assertEqual(ids, {self.s_visible.pk})


# ---------------------------------------------------------------------------
# PatientAdmin — main administrator
# ---------------------------------------------------------------------------

class PatientAdminSuperuserTest(TestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('pat_su', 'su@t.com', 'pass123')
        self.client.force_login(self.superuser)
        self.hidden = _patient('HADM01', 'مخفی ادمین', hidden=True)
        self.visible = _patient('HADM02', 'مشهود ادمین')

    def test_can_see_hidden_patient_in_list(self):
        resp = self.client.get('/admin/surgeries/patient/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'مخفی ادمین')

    def test_can_open_hidden_patient_change_page(self):
        resp = self.client.get(f'/admin/surgeries/patient/{self.hidden.pk}/change/')
        self.assertEqual(resp.status_code, 200)

    def test_can_search_hidden_patient(self):
        resp = self.client.get('/admin/surgeries/patient/', {'q': 'HADM01'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'مخفی ادمین')

    def test_can_filter_by_hidden(self):
        resp = self.client.get('/admin/surgeries/patient/', {'visibility': 'hidden'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'مخفی ادمین')
        self.assertNotContains(resp, 'مشهود ادمین')

    def test_bulk_hide_action(self):
        resp = self.client.post('/admin/surgeries/patient/', {
            'action': 'hide_selected_patients',
            '_selected_action': [str(self.visible.pk)],
        }, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.visible.refresh_from_db()
        self.assertTrue(self.visible.is_hidden)
        self.assertIsNotNone(self.visible.hidden_at)
        self.assertEqual(self.visible.hidden_by_id, self.superuser.pk)

    def test_bulk_unhide_action(self):
        resp = self.client.post('/admin/surgeries/patient/', {
            'action': 'unhide_selected_patients',
            '_selected_action': [str(self.hidden.pk)],
        }, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.hidden.refresh_from_db()
        self.assertFalse(self.hidden.is_hidden)
        self.assertIsNone(self.hidden.hidden_at)

    def test_hide_via_change_form_checkbox(self):
        resp = self.client.post(f'/admin/surgeries/patient/{self.visible.pk}/change/', {
            'full_name':    self.visible.full_name,
            'case_code':    self.visible.case_code,
            'phone_number': self.visible.phone_number,
            'national_id':  self.visible.national_id,
            'age':          '30',
            'gender':       'MALE',
            'description':  '',
            'is_hidden':    'on',
        }, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.visible.refresh_from_db()
        self.assertTrue(self.visible.is_hidden)
        self.assertEqual(self.visible.hidden_by_id, self.superuser.pk)


# ---------------------------------------------------------------------------
# PatientAdmin — normal staff user
# ---------------------------------------------------------------------------

class PatientAdminStaffTest(TestCase):

    def setUp(self):
        self.staff = _staff_with_patient_perms()
        self.client.force_login(self.staff)
        self.hidden  = _patient('SADM01', 'بیمار پنهان')
        self.hidden.is_hidden = True
        self.hidden.save(update_fields=['is_hidden'])
        self.visible = _patient('SADM02', 'بیمار عادی')

    def test_hidden_patient_absent_from_list(self):
        resp = self.client.get('/admin/surgeries/patient/')
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'بیمار پنهان')
        self.assertContains(resp, 'بیمار عادی')

    def test_hidden_patient_not_found_by_search(self):
        resp = self.client.get('/admin/surgeries/patient/', {'q': 'SADM01'})
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'بیمار پنهان')

    def test_direct_change_url_does_not_reveal_patient(self):
        # Django's own ModelAdmin convention for "not in this admin's
        # queryset" is a redirect to the admin index with a generic "does
        # not exist" message (the same thing that happens for an actually
        # deleted object) — not a raw 404, but it never reveals the hidden
        # Patient's data and never confirms it exists under a different id.
        resp = self.client.get(f'/admin/surgeries/patient/{self.hidden.pk}/change/')
        self.assertEqual(resp.status_code, 302)
        followed = self.client.get(f'/admin/surgeries/patient/{self.hidden.pk}/change/', follow=True)
        self.assertNotContains(followed, 'بیمار پنهان')
        self.assertNotContains(followed, self.hidden.case_code)

    def test_no_visibility_column_or_filter(self):
        resp = self.client.get('/admin/surgeries/patient/')
        self.assertNotContains(resp, 'visibility')

    def test_hide_action_not_available(self):
        resp = self.client.get('/admin/surgeries/patient/')
        self.assertNotContains(resp, 'hide_selected_patients')

    def test_forged_bulk_hide_action_denied(self):
        resp = self.client.post('/admin/surgeries/patient/', {
            'action': 'hide_selected_patients',
            '_selected_action': [str(self.visible.pk)],
        })
        # Django admin re-renders the changelist with an error for an
        # unavailable action rather than 500ing.
        self.assertIn(resp.status_code, (200, 400))
        self.visible.refresh_from_db()
        self.assertFalse(self.visible.is_hidden)

    def test_forged_is_hidden_post_on_own_patient_ignored(self):
        resp = self.client.post(f'/admin/surgeries/patient/{self.visible.pk}/change/', {
            'full_name':    self.visible.full_name,
            'case_code':    self.visible.case_code,
            'phone_number': self.visible.phone_number,
            'national_id':  self.visible.national_id,
            'age':          '30',
            'gender':       'MALE',
            'description':  '',
            'is_hidden':    'on',
        }, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.visible.refresh_from_db()
        self.assertFalse(self.visible.is_hidden)


# ---------------------------------------------------------------------------
# PatientViewSet (API)
# ---------------------------------------------------------------------------

class PatientApiVisibilityTest(APITestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('api_su', 'su@t.com', 'pass123')
        self.staff = _staff_with_patient_perms()
        self.hidden  = _patient('APIH01', 'مخفی API')
        self.hidden.is_hidden = True
        self.hidden.save(update_fields=['is_hidden'])
        self.visible = _patient('APIV01', 'مشهود API')

    def test_superuser_list_includes_hidden(self):
        self.client.force_authenticate(self.superuser)
        r = self.client.get(PATIENTS_URL)
        ids = {row['id'] for row in r.data['results']}
        self.assertIn(self.hidden.id, ids)

    def test_staff_list_excludes_hidden(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(PATIENTS_URL)
        ids = {row['id'] for row in r.data['results']}
        self.assertNotIn(self.hidden.id, ids)
        self.assertIn(self.visible.id, ids)

    def test_staff_search_does_not_find_hidden(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(PATIENTS_URL, {'search': 'مخفی API'})
        self.assertEqual(r.data['count'], 0)

    def test_staff_retrieve_hidden_returns_404(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{PATIENTS_URL}{self.hidden.id}/')
        self.assertEqual(r.status_code, 404)

    def test_superuser_retrieve_hidden_returns_200(self):
        self.client.force_authenticate(self.superuser)
        r = self.client.get(f'{PATIENTS_URL}{self.hidden.id}/')
        self.assertEqual(r.status_code, 200)

    def test_is_hidden_absent_for_staff(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{PATIENTS_URL}{self.visible.id}/')
        self.assertNotIn('is_hidden', r.data)

    def test_is_hidden_present_for_superuser(self):
        self.client.force_authenticate(self.superuser)
        r = self.client.get(f'{PATIENTS_URL}{self.visible.id}/')
        self.assertIn('is_hidden', r.data)

    def test_staff_forged_patch_is_hidden_ignored(self):
        self.client.force_authenticate(self.staff)
        r = self.client.patch(f'{PATIENTS_URL}{self.visible.id}/', {'is_hidden': True}, format='json')
        self.assertEqual(r.status_code, 200)
        self.visible.refresh_from_db()
        self.assertFalse(self.visible.is_hidden)

    def test_superuser_patch_is_hidden_applies_and_sets_audit(self):
        self.client.force_authenticate(self.superuser)
        r = self.client.patch(f'{PATIENTS_URL}{self.visible.id}/', {'is_hidden': True}, format='json')
        self.assertEqual(r.status_code, 200)
        self.visible.refresh_from_db()
        self.assertTrue(self.visible.is_hidden)
        self.assertEqual(self.visible.hidden_by_id, self.superuser.pk)
        self.assertIsNotNone(self.visible.hidden_at)

    def test_superuser_patch_is_hidden_false_restores(self):
        self.client.force_authenticate(self.superuser)
        self.client.patch(f'{PATIENTS_URL}{self.hidden.id}/', {'is_hidden': False}, format='json')
        self.hidden.refresh_from_db()
        self.assertFalse(self.hidden.is_hidden)
        self.assertIsNone(self.hidden.hidden_by_id)


# ---------------------------------------------------------------------------
# SurgeryHistory + hidden Patient
# ---------------------------------------------------------------------------

class SurgeryHistoryHiddenPatientTest(APITestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('sh_su', 'su@t.com', 'pass123')
        self.staff = _staff_with_patient_perms()
        self.patient = _patient('SHV01', 'بیمار جراحی')
        self.surgery = _surgery(self.patient)

    def test_hiding_patient_preserves_surgery_row(self):
        surgery_count_before = SurgeryHistory.objects.count()
        self.patient.hide(self.superuser)
        self.assertEqual(SurgeryHistory.objects.count(), surgery_count_before)
        self.surgery.refresh_from_db()
        self.assertEqual(self.surgery.patient_id, self.patient.id)

    def test_staff_list_excludes_surgery_of_hidden_patient(self):
        self.patient.hide(self.superuser)
        self.client.force_authenticate(self.staff)
        r = self.client.get(HISTORY_URL)
        ids = {row['id'] for row in r.data['results']}
        self.assertNotIn(self.surgery.id, ids)

    def test_superuser_list_includes_surgery_of_hidden_patient(self):
        self.patient.hide(self.superuser)
        self.client.force_authenticate(self.superuser)
        r = self.client.get(HISTORY_URL)
        ids = {row['id'] for row in r.data['results']}
        self.assertIn(self.surgery.id, ids)

    def test_staff_retrieve_surgery_of_hidden_patient_returns_404(self):
        self.patient.hide(self.superuser)
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{HISTORY_URL}{self.surgery.id}/')
        self.assertEqual(r.status_code, 404)

    def test_staff_search_does_not_reveal_hidden_patient_surgery(self):
        self.patient.hide(self.superuser)
        self.client.force_authenticate(self.staff)
        r = self.client.get(HISTORY_URL, {'search': 'بیمار جراحی'})
        self.assertEqual(r.data['count'], 0)

    def test_staff_cannot_create_surgery_for_hidden_patient(self):
        self.patient.hide(self.superuser)
        self.client.force_authenticate(self.staff)
        r = self.client.post(HISTORY_URL, {
            'patient': self.patient.id,
            'surgery_type': _surgery_type().id,
            'amount': '1000000',
            'surgery_date': '2025-07-01T10:00:00Z',
        }, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('patient', r.data)

    def test_restoring_patient_restores_full_visibility(self):
        self.patient.hide(self.superuser)
        self.patient.unhide()
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{HISTORY_URL}{self.surgery.id}/')
        self.assertEqual(r.status_code, 200)

    def test_no_duplicate_surgery_created_by_hide_restore_cycle(self):
        count_before = SurgeryHistory.objects.count()
        self.patient.hide(self.superuser)
        self.patient.unhide()
        self.assertEqual(SurgeryHistory.objects.count(), count_before)
        self.surgery.refresh_from_db()   # still the same row, unmodified

    def test_staff_admin_detail_view_returns_404_for_hidden_patient_surgery(self):
        from django.urls import reverse
        self.patient.hide(self.superuser)
        staff = _staff_with_patient_perms()
        self.client.force_login(staff)
        url = reverse('admin:surgeries_surgeryhistory_detail', args=[self.surgery.pk])
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 404)

    def test_superuser_admin_detail_view_returns_200_for_hidden_patient_surgery(self):
        from django.urls import reverse
        self.patient.hide(self.superuser)
        self.client.force_login(self.superuser)
        url = reverse('admin:surgeries_surgeryhistory_detail', args=[self.surgery.pk])
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)


class PatientAutocompleteVisibilityTest(TestCase):
    """Django admin's own /admin/autocomplete/ JSON endpoint, used by
    SurgeryHistoryAdmin's autocomplete_fields = ('patient', ...) — it
    delegates to PatientAdmin.get_queryset(), so this is really an
    integration check that the delegation actually happens."""

    def setUp(self):
        self.superuser = User.objects.create_superuser('ac_su', 'su@t.com', 'pass123')
        self.staff = _staff_with_patient_perms()
        self.staff.user_permissions.add(Permission.objects.get(codename='view_surgeryhistory'))
        self.hidden  = _patient('ACH01', 'بیمار اتوکامپلیت مخفی', hidden=True)
        self.visible = _patient('ACV01', 'بیمار اتوکامپلیت مشهود')

    def _autocomplete(self, term):
        return self.client.get('/admin/autocomplete/', {
            'app_label': 'surgeries', 'model_name': 'surgeryhistory',
            'field_name': 'patient', 'term': term,
        })

    def test_staff_cannot_find_hidden_patient_via_autocomplete(self):
        self.client.force_login(self.staff)
        resp = self._autocomplete('اتوکامپلیت')
        self.assertEqual(resp.status_code, 200)
        ids = {row['id'] for row in resp.json()['results']}
        self.assertNotIn(str(self.hidden.pk), ids)
        self.assertIn(str(self.visible.pk), ids)

    def test_superuser_can_find_hidden_patient_via_autocomplete(self):
        self.client.force_login(self.superuser)
        resp = self._autocomplete('اتوکامپلیت')
        self.assertEqual(resp.status_code, 200)
        ids = {row['id'] for row in resp.json()['results']}
        self.assertIn(str(self.hidden.pk), ids)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

class PatientExcelExportVisibilityTest(TestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('exp_su', 'su@t.com', 'pass123')
        self.hidden  = _patient('EXPH01', 'مخفی اکسل', hidden=True)
        self.visible = _patient('EXPV01', 'مشهود اکسل')

    def _open(self, content):
        import io
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_superuser_normal_export_excludes_hidden(self):
        self.client.force_login(self.superuser)
        r = self.client.get('/admin/surgeries/patient/export-excel/')
        self.assertEqual(r.status_code, 200)
        wb = self._open(r.content)
        all_vals = [
            ws.cell(row, col).value
            for ws in wb.worksheets
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertNotIn('مخفی اکسل', all_vals)
        self.assertIn('مشهود اکسل', all_vals)


class SurgeryHistoryExcelExportVisibilityTest(APITestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('sh_exp_su', 'su@t.com', 'pass123')
        self.client.force_authenticate(self.superuser)
        self.hidden_patient  = _patient('SHEH01', 'بیمار مخفی صادرات', hidden=True)
        self.visible_patient = _patient('SHEV01', 'بیمار مشهود صادرات')
        self.s_hidden  = _surgery(self.hidden_patient)
        self.s_visible = _surgery(self.visible_patient)

    def _open(self, content):
        import io
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_superuser_export_excludes_hidden_patient_surgery(self):
        r = self.client.get(HISTORY_URL, {'export': 'excel'})
        self.assertEqual(r.status_code, 200)
        wb = self._open(r.content)
        all_vals = [
            ws.cell(row, col).value
            for ws in wb.worksheets
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertNotIn('بیمار مخفی صادرات', all_vals)
        self.assertIn('بیمار مشهود صادرات', all_vals)
