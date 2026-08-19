"""Excel export tests for surgery endpoints (CLI-54).

Covers: /api/v2/surgeries/history/?export=excel
        /api/v2/surgeries/reports/profit/?export=excel
"""

import io
import datetime
from decimal import Decimal

from django.contrib.auth.models import Group, User
from rest_framework import status
from rest_framework.test import APITestCase

from finance.models import CenterCommissionIncome, IncomeStatus
from inventory.models import Product, ProductType
from surgeries.models import (
    Patient,
    PaymentStatus,
    SurgeryHistory,
    SurgeryStatus,
    SurgeryType,
    SurgeryUsedItem,
)

HISTORY_URL = '/api/v2/surgeries/history/'
PROFIT_URL  = '/api/v2/surgeries/reports/profit/'
XLSX_CT     = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _admin():
    u = User.objects.create_user(username=f'surg_adm_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='admin')
    u.groups.add(g)
    return u


def _finance_user():
    u = User.objects.create_user(username=f'fin_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='finance_user')
    u.groups.add(g)
    return u


def _surgery_type(name=None):
    import random
    uid = random.randint(100000, 999999)
    return SurgeryType.objects.create(
        name=name or f'عمل_{uid}', code=f'op_{uid}', base_rate=Decimal('1000000'),
    )


def _patient(name=None):
    import random
    uid = random.randint(100000, 999999)
    return Patient.objects.create(
        full_name=name or f'بیمار_{uid}',
        case_code=f'P-{uid}',
        phone_number='09000000000',
    )


def _surgery(patient, surgery_type, amount=Decimal('5000000'),
             surgery_date=None, surgery_status=SurgeryStatus.COMPLETED,
             payment_status=PaymentStatus.PAID):
    return SurgeryHistory.objects.create(
        patient=patient, surgery_type=surgery_type,
        amount=amount,
        surgery_date=surgery_date or datetime.datetime(2025, 6, 1, 10, 0),
        status=surgery_status,
        payment_status=payment_status,
    )


def _income(surgery, amount):
    return CenterCommissionIncome.objects.create(
        surgery=surgery, amount=amount,
        status=IncomeStatus.CONFIRMED,
        income_date=surgery.surgery_date,
    )


def _open_workbook(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


# ---------------------------------------------------------------------------
# Surgery History export
# ---------------------------------------------------------------------------

class SurgeryHistoryExcelExportTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        self.st  = _surgery_type('بینی')
        self.pat = _patient('رضا کریمی')
        self.s   = _surgery(self.pat, self.st, amount=Decimal('10000000'))
        self.s.center_commission_percent = Decimal('20')
        self.s.save(update_fields=['center_commission_percent'])
        _income(self.s, Decimal('2000000'))

    def _get(self, **params):
        return self.client.get(HISTORY_URL, {'export': 'excel', **params})

    def test_returns_xlsx_content_type(self):
        r = self._get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_content_disposition_has_filename(self):
        r = self._get()
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('.xlsx',      r['Content-Disposition'])

    def test_rtl_worksheet(self):
        r = self._get()
        ws = _open_workbook(r.content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_persian_headers(self):
        r = self._get()
        ws = _open_workbook(r.content).active
        headers = self._header_row(ws)
        self.assertIn('نام بیمار',             headers)
        self.assertIn('نوع عمل',               headers)
        self.assertIn('کمیسیون مرکز (تومان)', headers)
        self.assertIn('وضعیت عمل',             headers)

    def test_center_income_in_row(self):
        # Commission is computed live from the surgery's own configured
        # percentage (20% of amount=10,000,000 = 2,000,000) via
        # SurgeryFinanceService.calculate_center_commission — not read from
        # the separately-synced CenterCommissionIncome relation, which can
        # drift out of sync with the surgery's current amount/percent.
        r = self._get()
        ws = _open_workbook(r.content).active
        headers = self._header_row(ws)
        header_row_idx = self._header_row_index(ws)
        income_col = headers.index('کمیسیون مرکز (تومان)') + 1
        val = ws.cell(header_row_idx + 1, income_col).value
        self.assertAlmostEqual(float(val), 2000000.0)

    @staticmethod
    def _header_row_index(ws):
        for row in range(1, ws.max_row + 1):
            if any(ws.cell(row, c).value == 'نام بیمار' for c in range(1, ws.max_column + 1)):
                return row
        raise AssertionError('header row not found')

    def _header_row(self, ws):
        row = self._header_row_index(ws)
        return [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]

    def _patient_name_column_values(self, ws):
        header_row = self._header_row_index(ws)
        headers    = self._header_row(ws)
        name_col   = headers.index('نام بیمار') + 1
        return [
            ws.cell(row, name_col).value
            for row in range(header_row + 1, ws.max_row + 1)
            if ws.cell(row, name_col).value
        ]

    def test_surgery_type_filter(self):
        st2 = _surgery_type('گوش')
        s2  = _surgery(_patient('نفر دوم'), st2)
        r   = self._get(surgery_type=self.st.id)
        ws  = _open_workbook(r.content).active
        names = self._patient_name_column_values(ws)
        self.assertIn('رضا کریمی', names)
        self.assertNotIn('نفر دوم', names)

    def test_payment_status_filter(self):
        s_pending = _surgery(_patient(), self.st, payment_status=PaymentStatus.PENDING)
        r = self._get(payment_status='PAID')
        ws = _open_workbook(r.content).active
        # Only PAID surgeries (setUp surgery)
        data_rows = self._patient_name_column_values(ws)
        self.assertIn('رضا کریمی', data_rows)
        self.assertEqual(len(data_rows), 1)

    def test_no_pagination(self):
        for i in range(5):
            _surgery(_patient(), self.st)
        r = self._get(page_size='2')
        ws = _open_workbook(r.content).active
        rows = self._patient_name_column_values(ws)
        self.assertGreaterEqual(len(rows), 6)

    def test_date_filtered_export_contains_only_matching_surgeries(self):
        # setUp's self.s is 2025-06-01. Add one inside and one outside a
        # June 2025 date range, then confirm the export reflects exactly
        # the same date-range filter as the (JS-driven) list page.
        s_in = _surgery(
            _patient('در بازه'), self.st,
            surgery_date=datetime.datetime(2025, 6, 15, 10, 0),
        )
        s_out = _surgery(
            _patient('خارج بازه'), self.st,
            surgery_date=datetime.datetime(2025, 1, 1, 10, 0),
        )
        r = self._get(surgery_date_from='2025-06-01', surgery_date_to='2025-06-30')
        self.assertEqual(r.status_code, 200)
        ws = _open_workbook(r.content).active
        names = self._patient_name_column_values(ws)
        self.assertIn('رضا کریمی', names)   # setUp surgery, 2025-06-01 (inclusive boundary)
        self.assertIn('در بازه',   names)
        self.assertNotIn('خارج بازه', names)

    def test_reversed_date_range_export_returns_error_not_500(self):
        r = self._get(surgery_date_from='2025-06-30', surgery_date_to='2025-06-01')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_unauthenticated_returns_401(self):
        self.client.force_authenticate(None)
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_401_UNAUTHORIZED)


# ---------------------------------------------------------------------------
# Surgery Profit Report export
# ---------------------------------------------------------------------------

class SurgeryProfitReportExcelTest(APITestCase):

    def setUp(self):
        self.user = _finance_user()
        self.client.force_authenticate(self.user)
        self.st  = _surgery_type()
        self.pat = _patient()
        self.s   = _surgery(self.pat, self.st, amount=Decimal('5000000'))
        _income(self.s, Decimal('1000000'))

    def _get(self, **params):
        return self.client.get(PROFIT_URL, {'export': 'excel', **params})

    def test_group_by_surgery_returns_xlsx(self):
        r = self._get(group_by='surgery')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_group_by_surgery_persian_headers(self):
        r = self._get(group_by='surgery')
        ws = _open_workbook(r.content).active
        all_vals = [ws.cell(row, col).value for row in range(1, ws.max_row + 1)
                    for col in range(1, ws.max_column + 1) if ws.cell(row, col).value]
        self.assertIn('بیمار',          all_vals)
        self.assertIn('درآمد مرکز',     all_vals)
        self.assertIn('سود تقریبی',     all_vals)

    def test_group_by_surgery_type_returns_xlsx(self):
        r = self._get(group_by='surgery_type')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_group_by_surgery_type_headers(self):
        r = self._get(group_by='surgery_type')
        ws = _open_workbook(r.content).active
        all_vals = [ws.cell(row, col).value for row in range(1, ws.max_row + 1)
                    for col in range(1, ws.max_column + 1) if ws.cell(row, col).value]
        self.assertIn('نوع عمل',               all_vals)
        self.assertIn('مجموع سود تقریبی',      all_vals)
        self.assertIn('میانگین سود هر عمل',    all_vals)

    def test_date_filter_applied(self):
        # Surgery in 2024 should be excluded when start_date=2025-01-01
        _surgery(_patient(), self.st, surgery_date=datetime.datetime(2024, 1, 1, 10, 0))
        r   = self._get(group_by='surgery', start_date='2025-01-01')
        ws  = _open_workbook(r.content).active
        # Find header row (contains 'بیمار')
        header_row = None
        for row in range(1, ws.max_row + 1):
            if any(ws.cell(row, c).value == 'بیمار' for c in range(1, ws.max_column + 1)):
                header_row = row
                break
        self.assertIsNotNone(header_row)
        data_names = [ws.cell(row, 1).value for row in range(header_row + 1, ws.max_row + 1)
                      if ws.cell(row, 1).value]
        # Only the setUp surgery (2025) should appear — not the 2024 one
        self.assertEqual(len(data_names), 1)

    def test_finance_user_can_export(self):
        r = self._get(group_by='surgery')
        self.assertEqual(r.status_code, 200)

    def test_unauthorized_user_gets_403(self):
        u = User.objects.create_user(username='plain_user', password='x')
        self.client.force_authenticate(u)
        r = self._get(group_by='surgery')
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)
