"""Tests for Patient age/gender and their automatic display on the
Surgery form/detail page.

Existing architecture found before this change: no age/birth_date/gender
field existed anywhere on Patient. SurgeryHistorySerializer already had a
`patient_age` stub (`get_patient_age`) that checked for a Patient `.age`
attribute first, before falling back to birth_date/date_of_birth — written
in anticipation of exactly this field, so adding `Patient.age` as a direct
stored field (not birth-date-derived) made that stub start working with no
further serializer change for age. No SurgeryHistory snapshot fields were
added — Patient remains the single canonical source for age/gender,
per the project's own stated preference, since nothing existing depends on
preserving age/gender as of a past Surgery date.
"""

import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.test import TestCase
from rest_framework.test import APITestCase

from surgeries.forms import PatientAdminForm
from surgeries.models import Gender, Patient, SurgeryHistory, SurgeryType

User = get_user_model()

PATIENTS_URL = '/api/v2/surgeries/patients/'
HISTORY_URL  = '/api/v2/surgeries/history/'


def _patient(case_code='PAT001', full_name='بیمار تست', age=30, gender=Gender.MALE, hidden=False):
    p = Patient.objects.create(
        full_name=full_name, case_code=case_code,
        phone_number='09120000000', national_id='5555555555',
        age=age, gender=gender,
    )
    if hidden:
        p.is_hidden = True
        p.save(update_fields=['is_hidden'])
    return p


def _surgery_type():
    n = SurgeryType.objects.count()
    return SurgeryType.objects.create(name=f'عمومی_{n}', code=f'gen_{n}', base_rate=Decimal('1000000'))


def _surgery(patient, surgery_type=None):
    return SurgeryHistory.objects.create(
        patient=patient,
        surgery_type=surgery_type or _surgery_type(),
        amount=Decimal('10000000'),
        surgery_date=datetime.datetime(2025, 6, 1, 10, 0, tzinfo=datetime.timezone.utc),
    )


def _staff_with_patient_perms():
    staff = User.objects.create_user(
        f'ag_staff_{User.objects.count()}', 'staff@t.com', 'pass123', is_staff=True,
    )
    for codename in ('view_patient', 'add_patient', 'change_patient', 'view_surgeryhistory', 'add_surgeryhistory'):
        staff.user_permissions.add(Permission.objects.get(codename=codename))
    return staff


def _valid_patient_form_data(**overrides):
    data = {
        'full_name':    'بیمار جدید',
        'case_code':    f'NEWPAT-{Patient.objects.count()}',
        'national_id':  '1234567890',
        'phone_number': '09120001111',
        'age':          '40',
        'gender':       Gender.MALE,
        'description':  '',
    }
    data.update(overrides)
    return data


# ---------------------------------------------------------------------------
# Patient model / form
# ---------------------------------------------------------------------------

class PatientAgeGenderModelTest(TestCase):

    def test_gender_choices_are_only_male_and_female(self):
        values = {choice[0] for choice in Gender.choices}
        self.assertEqual(values, {'MALE', 'FEMALE'})

    def test_create_patient_with_valid_age_and_male_gender(self):
        p = _patient('AG001', age=25, gender=Gender.MALE)
        self.assertEqual(p.age, 25)
        self.assertEqual(p.gender, 'MALE')
        self.assertEqual(p.get_gender_display(), 'مرد')

    def test_create_patient_with_valid_age_and_female_gender(self):
        p = _patient('AG002', age=25, gender=Gender.FEMALE)
        self.assertEqual(p.gender, 'FEMALE')
        self.assertEqual(p.get_gender_display(), 'زن')

    def test_invalid_gender_rejected_by_full_clean(self):
        p = Patient(full_name='x', case_code='AG003', phone_number='0912', age=30, gender='OTHER')
        with self.assertRaises(ValidationError):
            p.full_clean()

    def test_negative_age_rejected_by_full_clean(self):
        p = Patient(full_name='x', case_code='AG004', phone_number='0912', age=-1, gender=Gender.MALE)
        with self.assertRaises(ValidationError):
            p.full_clean()

    def test_age_above_maximum_rejected_by_full_clean(self):
        p = Patient(full_name='x', case_code='AG005', phone_number='0912', age=131, gender=Gender.MALE)
        with self.assertRaises(ValidationError):
            p.full_clean()

    def test_existing_incomplete_historical_patient_remains_readable(self):
        # Simulates a pre-existing row with no age/gender ever entered.
        p = Patient.objects.create(
            full_name='بیمار قدیمی', case_code='OLD001', phone_number='0912',
        )
        p.refresh_from_db()
        self.assertIsNone(p.age)
        self.assertIsNone(p.gender)
        self.assertIsNone(p.get_gender_display())


class PatientAdminFormTest(TestCase):

    def test_age_and_gender_required_on_admin_form(self):
        form = PatientAdminForm(data=_valid_patient_form_data(age='', gender=''))
        self.assertFalse(form.is_valid())
        self.assertIn('age', form.errors)
        self.assertIn('gender', form.errors)

    def test_valid_admin_form_is_accepted(self):
        form = PatientAdminForm(data=_valid_patient_form_data())
        self.assertTrue(form.is_valid(), form.errors)

    def test_admin_form_rejects_invalid_gender_choice(self):
        form = PatientAdminForm(data=_valid_patient_form_data(gender='OTHER'))
        self.assertFalse(form.is_valid())
        self.assertIn('gender', form.errors)

    def test_admin_form_rejects_negative_age(self):
        form = PatientAdminForm(data=_valid_patient_form_data(age='-5'))
        self.assertFalse(form.is_valid())
        self.assertIn('age', form.errors)

    def test_admin_form_rejects_age_above_maximum(self):
        form = PatientAdminForm(data=_valid_patient_form_data(age='131'))
        self.assertFalse(form.is_valid())
        self.assertIn('age', form.errors)


class PatientAdminChangeFormTest(TestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('agf_su', 'su@t.com', 'pass123')
        self.client.force_login(self.superuser)

    def test_add_form_contains_age_and_gender_fields(self):
        resp = self.client.get('/admin/surgeries/patient/add/')
        self.assertContains(resp, 'id="id_age"')
        self.assertContains(resp, 'id="id_gender"')
        self.assertContains(resp, 'سن بیمار')
        self.assertContains(resp, 'جنسیت')

    def test_add_form_gender_options_are_only_male_female(self):
        resp = self.client.get('/admin/surgeries/patient/add/')
        content = resp.content.decode('utf-8')
        self.assertIn('مرد', content)
        self.assertIn('زن', content)
        self.assertNotIn('سایر', content)

    def test_create_patient_via_admin_persists_age_gender(self):
        resp = self.client.post('/admin/surgeries/patient/add/', _valid_patient_form_data(
            case_code='ADMNEW01',
        ))
        self.assertIn(resp.status_code, (200, 302))
        p = Patient.objects.get(case_code='ADMNEW01')
        self.assertEqual(p.age, 40)
        self.assertEqual(p.gender, 'MALE')

    def test_change_form_persists_age_and_gender(self):
        patient = _patient('ADMCHG01', age=22, gender=Gender.FEMALE)
        resp = self.client.post(f'/admin/surgeries/patient/{patient.pk}/change/', _valid_patient_form_data(
            case_code='ADMCHG01', age='55', gender=Gender.MALE,
        ))
        self.assertIn(resp.status_code, (200, 302))
        patient.refresh_from_db()
        self.assertEqual(patient.age, 55)
        self.assertEqual(patient.gender, 'MALE')

    def test_change_form_rejects_invalid_gender(self):
        patient = _patient('ADMCHG02', gender=Gender.MALE)
        resp = self.client.post(f'/admin/surgeries/patient/{patient.pk}/change/', _valid_patient_form_data(
            case_code='ADMCHG02', gender='OTHER',
        ))
        self.assertEqual(resp.status_code, 200)   # re-rendered with errors, not redirected
        self.assertContains(resp, 'class="errorlist"')
        patient.refresh_from_db()
        self.assertEqual(patient.gender, 'MALE')   # unchanged — invalid submission never saved


# ---------------------------------------------------------------------------
# Patient API
# ---------------------------------------------------------------------------

class PatientApiAgeGenderTest(APITestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('agapi_su', 'su@t.com', 'pass123')
        self.staff = _staff_with_patient_perms()
        self.patient = _patient('API001', age=35, gender=Gender.MALE)
        self.hidden  = _patient('API002', age=40, gender=Gender.FEMALE, hidden=True)

    def test_age_is_numeric_in_response(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{PATIENTS_URL}{self.patient.id}/')
        self.assertEqual(r.status_code, 200)
        self.assertIsInstance(r.data['age'], int)
        self.assertEqual(r.data['age'], 35)

    def test_persian_gender_display_returned(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{PATIENTS_URL}{self.patient.id}/')
        self.assertEqual(r.data['gender'], 'MALE')
        self.assertEqual(r.data['gender_display'], 'مرد')

    def test_female_gender_display(self):
        female = _patient('API003', age=28, gender=Gender.FEMALE)
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{PATIENTS_URL}{female.id}/')
        self.assertEqual(r.data['gender_display'], 'زن')

    def test_list_serializer_includes_age_gender(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(PATIENTS_URL)
        row = next(x for x in r.data['results'] if x['id'] == self.patient.id)
        self.assertEqual(row['age'], 35)
        self.assertEqual(row['gender_display'], 'مرد')

    def test_hidden_patient_age_gender_not_retrievable_by_unauthorized_user(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(f'{PATIENTS_URL}{self.hidden.id}/')
        self.assertEqual(r.status_code, 404)

    def test_hidden_patient_age_gender_not_in_staff_list(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(PATIENTS_URL)
        ids = {row['id'] for row in r.data['results']}
        self.assertNotIn(self.hidden.id, ids)

    def test_superuser_can_retrieve_hidden_patient_age_gender(self):
        self.client.force_authenticate(self.superuser)
        r = self.client.get(f'{PATIENTS_URL}{self.hidden.id}/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['age'], 40)
        self.assertEqual(r.data['gender_display'], 'زن')

    def test_create_patient_with_invalid_gender_rejected(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(PATIENTS_URL, {
            'full_name': 'تست', 'case_code': 'APIBAD01', 'phone_number': '0912',
            'age': 30, 'gender': 'OTHER',
        }, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('gender', r.data)

    def test_create_patient_with_negative_age_rejected(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(PATIENTS_URL, {
            'full_name': 'تست', 'case_code': 'APIBAD02', 'phone_number': '0912',
            'age': -3, 'gender': 'MALE',
        }, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('age', r.data)

    def test_create_patient_with_age_above_maximum_rejected(self):
        self.client.force_authenticate(self.staff)
        r = self.client.post(PATIENTS_URL, {
            'full_name': 'تست', 'case_code': 'APIBAD03', 'phone_number': '0912',
            'age': 200, 'gender': 'MALE',
        }, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('age', r.data)


# ---------------------------------------------------------------------------
# Surgery form / API — automatic Patient metadata
# ---------------------------------------------------------------------------

class SurgeryPatientMetadataApiTest(APITestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('shag_su', 'su@t.com', 'pass123')
        self.client.force_authenticate(self.superuser)
        self.patient_a = _patient('SHAG001', full_name='بیمار الف', age=45, gender=Gender.MALE)
        self.patient_b = _patient('SHAG002', full_name='بیمار ب', age=29, gender=Gender.FEMALE)
        self.surgery = _surgery(self.patient_a)

    def test_surgery_detail_returns_patient_age_and_gender(self):
        r = self.client.get(f'{HISTORY_URL}{self.surgery.id}/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['patient_age'], 45)
        self.assertEqual(r.data['patient_gender'], 'MALE')
        self.assertEqual(r.data['patient_gender_display'], 'مرد')

    def test_surgery_list_includes_patient_age_gender(self):
        r = self.client.get(HISTORY_URL)
        row = next(x for x in r.data['results'] if x['id'] == self.surgery.id)
        self.assertEqual(row['patient_age'], 45)
        self.assertEqual(row['patient_gender_display'], 'مرد')

    def test_different_patients_report_different_metadata(self):
        # This is exactly what the Surgery form's JS re-fetches on Patient
        # selector change — confirms the underlying data genuinely differs
        # per Patient, the premise the whole auto-population feature relies on.
        ra = self.client.get(f'{PATIENTS_URL}{self.patient_a.id}/')
        rb = self.client.get(f'{PATIENTS_URL}{self.patient_b.id}/')
        self.assertNotEqual(ra.data['age'], rb.data['age'])
        self.assertNotEqual(ra.data['gender_display'], rb.data['gender_display'])

    def test_client_submitted_patient_age_gender_are_ignored_on_create(self):
        # SurgeryHistory has no age/gender fields of its own — a forged
        # patient_age/patient_gender in the request body must be silently
        # ignored (they are read_only), and the response must reflect the
        # real Patient record, not the submitted values.
        r = self.client.post(HISTORY_URL, {
            'patient': self.patient_b.id,
            'surgery_type': _surgery_type().id,
            'amount': '2000000',
            'surgery_date': '2025-07-01T10:00:00Z',
            'patient_age': 999,
            'patient_gender': 'MALE',
            'patient_gender_display': 'فروخته‌شده',
        }, format='json')
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.data['patient_age'], 29)
        self.assertEqual(r.data['patient_gender_display'], 'زن')
        surgery = SurgeryHistory.objects.get(pk=r.data['id'])
        self.assertEqual(surgery.patient_id, self.patient_b.id)

    def test_hidden_patient_cannot_be_selected_for_new_surgery(self):
        staff = _staff_with_patient_perms()
        hidden_patient = _patient('SHAG003', hidden=True)
        self.client.force_authenticate(staff)
        r = self.client.post(HISTORY_URL, {
            'patient': hidden_patient.id,
            'surgery_type': _surgery_type().id,
            'amount': '1000000',
            'surgery_date': '2025-07-01T10:00:00Z',
        }, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('patient', r.data)


# ---------------------------------------------------------------------------
# Excel exports
# ---------------------------------------------------------------------------

class PatientExcelAgeGenderExportTest(TestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('agexp_su', 'su@t.com', 'pass123')
        self.client.force_login(self.superuser)
        self.visible = _patient('EXAG01', full_name='قابل مشاهده اکسل', age=50, gender=Gender.FEMALE)
        self.hidden  = _patient('EXAG02', full_name='مخفی اکسل', age=60, gender=Gender.MALE, hidden=True)

    def _open(self, content):
        import io
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_export_includes_readable_age_and_gender(self):
        r = self.client.get('/admin/surgeries/patient/export-excel/')
        self.assertEqual(r.status_code, 200)
        wb = self._open(r.content)
        all_vals = [
            ws.cell(row, col).value
            for ws in wb.worksheets
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertIn(50, all_vals)
        self.assertIn('زن', all_vals)
        self.assertNotIn('MALE', all_vals)
        self.assertNotIn('FEMALE', all_vals)

    def test_hidden_patient_age_gender_excluded_from_export(self):
        r = self.client.get('/admin/surgeries/patient/export-excel/')
        wb = self._open(r.content)
        all_vals = [
            ws.cell(row, col).value
            for ws in wb.worksheets
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertNotIn('مخفی اکسل', all_vals)
        self.assertNotIn(60, all_vals)


class SurgeryExcelAgeGenderExportTest(APITestCase):

    def setUp(self):
        self.superuser = User.objects.create_superuser('shexp_su', 'su@t.com', 'pass123')
        self.client.force_authenticate(self.superuser)
        self.visible_patient = _patient('SHEXAG01', full_name='بیمار صادرات مشهود', age=33, gender=Gender.MALE)
        self.hidden_patient  = _patient('SHEXAG02', full_name='بیمار صادرات مخفی', age=44, gender=Gender.FEMALE, hidden=True)
        self.s_visible = _surgery(self.visible_patient)
        self.s_hidden  = _surgery(self.hidden_patient)

    def _open(self, content):
        import io
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content))

    def test_export_includes_readable_patient_age_gender(self):
        r = self.client.get(HISTORY_URL, {'export': 'excel'})
        self.assertEqual(r.status_code, 200)
        wb = self._open(r.content)
        all_vals = [
            ws.cell(row, col).value
            for ws in wb.worksheets
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertIn(33, all_vals)
        self.assertIn('مرد', all_vals)

    def test_hidden_patient_surgery_and_metadata_excluded_from_export(self):
        r = self.client.get(HISTORY_URL, {'export': 'excel'})
        wb = self._open(r.content)
        all_vals = [
            ws.cell(row, col).value
            for ws in wb.worksheets
            for row in range(1, ws.max_row + 1)
            for col in range(1, ws.max_column + 1)
        ]
        self.assertNotIn('بیمار صادرات مخفی', all_vals)
        self.assertNotIn(44, all_vals)
