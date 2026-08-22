"""Unit tests for the shared Excel export service (common/excel.py).

These test build_workbook/build_excel/ExcelColumn/ExcelSheet directly,
independent of any particular app's ViewSet or ModelAdmin — the app-level
export tests (surgeries/inventory/employees/contacts) cover integration
through the real endpoints instead.
"""

import io
from datetime import date
from decimal import Decimal

from django.test import SimpleTestCase

from common.excel import (
    DEFAULT_EMPTY_MESSAGE,
    ExcelColumn,
    ExcelSheet,
    build_excel,
    build_workbook,
    describe_ordering,
    make_export_filename,
)


def _open(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


class BuildWorkbookTest(SimpleTestCase):

    def test_decimal_values_stay_numeric_and_precise(self):
        # openpyxl has no native Decimal cell type — numbers always read back
        # as float/int — but writing a Decimal directly (instead of calling
        # float() ourselves first) avoids any float rounding error creeping
        # in on the write side. Verify the cell is numeric (not a
        # money-with-currency-suffix string) and precise to the cent.
        columns = [ExcelColumn(key='amount', label='مبلغ', data_type='money')]
        rows = [{'amount': Decimal('1234567.89')}]
        content = build_excel(columns=columns, rows=rows)
        ws = _open(content).active
        header_row = next(
            r for r in range(1, ws.max_row + 1)
            if ws.cell(r, 1).value == 'مبلغ'
        )
        value = ws.cell(header_row + 1, 1).value
        self.assertIsInstance(value, (int, float))
        self.assertEqual(round(float(value), 2), 1234567.89)

    def test_zero_and_none_are_distinguishable(self):
        columns = [ExcelColumn(key='n', label='عدد', data_type='integer')]
        rows = [{'n': 0}, {'n': None}]
        content = build_excel(columns=columns, rows=rows)
        ws = _open(content).active
        header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'عدد')
        self.assertEqual(ws.cell(header_row + 1, 1).value, 0)
        self.assertIsNone(ws.cell(header_row + 2, 1).value)

    def test_none_date_renders_as_dash(self):
        columns = [ExcelColumn(key='d', label='تاریخ', data_type='date')]
        rows = [{'d': None}, {'d': date(2025, 6, 1)}]
        content = build_excel(columns=columns, rows=rows)
        ws = _open(content).active
        header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'تاریخ')
        self.assertEqual(ws.cell(header_row + 1, 1).value, '—')
        self.assertNotEqual(ws.cell(header_row + 2, 1).value, '—')

    def test_empty_rows_show_empty_message_not_error(self):
        columns = [ExcelColumn(key='x', label='ستون', data_type='text')]
        content = build_excel(columns=columns, rows=[])
        ws = _open(content).active
        all_vals = [
            ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
        ]
        self.assertIn(DEFAULT_EMPTY_MESSAGE, all_vals)

    def test_custom_empty_message(self):
        sheet = ExcelSheet(
            name='گزارش', columns=[ExcelColumn(key='x', label='ستون')],
            rows=[], empty_message='چیزی یافت نشد!',
        )
        content = build_workbook([sheet])
        ws = _open(content).active
        all_vals = [
            ws.cell(r, c).value
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
        ]
        self.assertIn('چیزی یافت نشد!', all_vals)

    def test_rtl_sheet_view(self):
        content = build_excel(columns=[ExcelColumn(key='x', label='ستون')], rows=[])
        ws = _open(content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_multi_sheet_workbook_has_both_sheets(self):
        s1 = ExcelSheet(name='برگه یک', columns=[ExcelColumn(key='a', label='الف')], rows=[{'a': 1}])
        s2 = ExcelSheet(name='برگه دو', columns=[ExcelColumn(key='b', label='ب')], rows=[{'b': 2}])
        content = build_workbook([s1, s2])
        wb = _open(content)
        self.assertEqual(wb.sheetnames, ['برگه یک', 'برگه دو'])

    def test_report_title_and_meta_rows_appear(self):
        content = build_excel(
            columns=[ExcelColumn(key='x', label='ستون')], rows=[],
            report_title='گزارش تست', meta_rows=[('فیلتر', 'فعال')],
        )
        ws = _open(content).active
        all_vals = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        self.assertIn('گزارش تست', all_vals)
        self.assertTrue(any(v and 'فیلتر' in str(v) for v in all_vals))

    def test_freeze_panes_and_autofilter_set(self):
        content = build_excel(columns=[ExcelColumn(key='x', label='ستون')], rows=[{'x': 1}])
        ws = _open(content).active
        self.assertIsNotNone(ws.freeze_panes)
        self.assertIsNotNone(ws.auto_filter.ref)

    def test_legacy_tuple_columns_still_supported(self):
        # Backward compatibility with the original (header, key) convention.
        content = build_excel(columns=[('نام', 'name')], rows=[{'name': 'علی'}])
        ws = _open(content).active
        all_vals = [ws.cell(r, c).value for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)]
        self.assertIn('نام', all_vals)
        self.assertIn('علی', all_vals)


class SheetNameSanitizationTest(SimpleTestCase):

    def test_long_and_duplicate_sheet_names_are_handled(self):
        long_name = 'ا' * 50
        s1 = ExcelSheet(name=long_name, columns=[ExcelColumn(key='a', label='a')], rows=[])
        s2 = ExcelSheet(name=long_name, columns=[ExcelColumn(key='b', label='b')], rows=[])
        content = build_workbook([s1, s2])
        wb = _open(content)
        self.assertEqual(len(wb.sheetnames), 2)
        for name in wb.sheetnames:
            self.assertLessEqual(len(name), 31)
        self.assertEqual(len(set(wb.sheetnames)), 2)


class MakeExportFilenameTest(SimpleTestCase):

    def test_filename_has_prefix_and_xlsx_extension(self):
        name = make_export_filename('surgery-list')
        self.assertTrue(name.startswith('surgery-list-'))
        self.assertTrue(name.endswith('.xlsx'))

    def test_unsafe_characters_in_prefix_are_sanitized(self):
        name = make_export_filename('a/b c?d')
        self.assertNotIn('/', name)
        self.assertNotIn(' ', name)
        self.assertNotIn('?', name)


class DescribeOrderingTest(SimpleTestCase):

    def test_none_returns_none(self):
        self.assertIsNone(describe_ordering(None, {}))

    def test_ascending_and_descending_labels(self):
        labels = {'name': 'نام', 'created_at': 'تاریخ ثبت'}
        self.assertEqual(describe_ordering('name', labels), 'نام (صعودی)')
        self.assertEqual(describe_ordering('-created_at', labels), 'تاریخ ثبت (نزولی)')

    def test_unknown_field_falls_back_to_raw_name(self):
        self.assertEqual(describe_ordering('mystery_field', {}), 'mystery_field (صعودی)')
