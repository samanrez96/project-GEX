"""Excel export tests for inventory endpoints (CLI-54).

Covers: /api/v1/inventory/products/?export=excel
        /api/v1/inventory/purchases/?export=excel
        /api/v1/inventory/reports/cost/?export=excel
"""

import io
from decimal import Decimal

from django.contrib.auth.models import Group, User
from rest_framework import status
from rest_framework.test import APITestCase

from inventory.models import (
    Product,
    ProductCategory,
    ProductType,
    Purchase,
    PurchaseItem,
    PurchaseStatus,
    Vendor,
)

PRODUCTS_URL  = '/api/v1/inventory/products/'
PURCHASES_URL = '/api/v1/inventory/purchases/'
COST_REPORT_URL = '/api/v1/inventory/reports/cost/'

XLSX_CT = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _admin():
    u = User.objects.create_user(username=f'adm_{User.objects.count()}', password='x')
    g, _ = Group.objects.get_or_create(name='admin')
    u.groups.add(g)
    return u


def _product(name=None, purchase_price=Decimal('10000'), is_active=True):
    import random
    uid = random.randint(100000, 999999)
    return Product.objects.create(
        name=name or f'Prod_{uid}',
        internal_code=f'C-{uid}',
        product_type=ProductType.MEDICINE,
        unit='عدد',
        purchase_price=purchase_price,
        is_active=is_active,
    )


def _vendor(name=None):
    import random
    return Vendor.objects.create(name=name or f'Vendor_{random.randint(1000, 9999)}')


def _purchase(vendor, status=PurchaseStatus.CONFIRMED):
    return Purchase.objects.create(vendor=vendor, status=status)


def _purchase_item(purchase, product, qty=Decimal('2'), unit_price=Decimal('5000')):
    return PurchaseItem.objects.create(
        purchase=purchase, product=product, quantity=qty, unit_price=unit_price,
    )


def _open_workbook(content: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


def _header_row_index(ws, *, contains):
    """Row index (1-based) of the header row — identified by a known label,
    since build_workbook() prefixes a title/meta/generated-line block whose
    length varies per report, so the header is never reliably row 1."""
    for row in range(1, ws.max_row + 1):
        if any(ws.cell(row, c).value == contains for c in range(1, ws.max_column + 1)):
            return row
    raise AssertionError(f'header row containing {contains!r} not found')


def _header_labels(ws, header_row):
    return [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]


def _column_values(ws, header_row, col_idx):
    return [
        ws.cell(row, col_idx).value
        for row in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row, col_idx).value not in (None, '')
    ]


# ---------------------------------------------------------------------------
# Product export
# ---------------------------------------------------------------------------

class ProductExcelExportTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        self.p1 = _product('محصول الف', is_active=True)
        self.p2 = _product('محصول ب',   is_active=False)

    def _get(self, **params):
        return self.client.get(PRODUCTS_URL, {'export': 'excel', **params})

    def test_returns_xlsx_content_type(self):
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_200_OK)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_content_disposition_has_filename(self):
        r = self._get()
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertIn('.xlsx', r['Content-Disposition'])

    def test_workbook_is_openable(self):
        r = self._get()
        wb = _open_workbook(r.content)
        self.assertIsNotNone(wb)

    def test_worksheet_is_rtl(self):
        r = self._get(is_active='all')
        ws = _open_workbook(r.content).active
        self.assertTrue(ws.sheet_view.rightToLeft)

    def test_persian_headers_present(self):
        r = self._get(is_active='all')
        ws = _open_workbook(r.content).active
        header_row = _header_row_index(ws, contains='نام محصول')
        header_labels = _header_labels(ws, header_row)
        self.assertIn('نام محصول', header_labels)
        self.assertIn('کد داخلی',  header_labels)
        self.assertIn('موجودی فعلی', header_labels)

    def _name_column_values(self, ws):
        header_row = _header_row_index(ws, contains='نام محصول')
        name_col   = _header_labels(ws, header_row).index('نام محصول') + 1
        return _column_values(ws, header_row, name_col)

    def test_active_filter_applied(self):
        # Default: only active products
        r = self._get()
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('محصول الف', names)
        self.assertNotIn('محصول ب', names)

    def test_is_active_all_returns_all(self):
        r = self._get(is_active='all')
        ws = _open_workbook(r.content).active
        names = self._name_column_values(ws)
        self.assertIn('محصول الف', names)
        self.assertIn('محصول ب',   names)

    def test_no_pagination_all_rows_exported(self):
        # Create 5 products; even with page_size=2 all should appear in Excel
        for i in range(5):
            _product(f'Bulk_{i}')
        r = self._get(is_active='all', page_size='2')
        ws = _open_workbook(r.content).active
        # 2 from setUp + 5 new = 7; Excel should have all
        data_rows = self._name_column_values(ws)
        self.assertGreaterEqual(len(data_rows), 7)

    def test_unauthenticated_returns_401(self):
        self.client.force_authenticate(None)
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_401_UNAUTHORIZED)


# ---------------------------------------------------------------------------
# Purchase export
# ---------------------------------------------------------------------------

class PurchaseExcelExportTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        self.vendor = _vendor('تامین‌کننده آ')
        self.prod   = _product('کالا یک', purchase_price=Decimal('100'))
        self.purchase = _purchase(self.vendor, status=PurchaseStatus.CONFIRMED)
        _purchase_item(self.purchase, self.prod, qty=Decimal('3'), unit_price=Decimal('100'))

    def _get(self, **params):
        return self.client.get(PURCHASES_URL, {'export': 'excel', **params})

    def test_returns_xlsx_content_type(self):
        r = self._get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    # Purchase export is a two-sheet workbook: "خلاصه خریدها" (one row per
    # Purchase, the active/first sheet) and "اقلام خرید" (one row per
    # PurchaseItem) — item-level assertions (product, quantity, unit price)
    # must read the second sheet by name.

    def _items_sheet(self, content):
        return _open_workbook(content)['اقلام خرید']

    def test_persian_headers(self):
        r = self._get()
        ws = self._items_sheet(r.content)
        header_row = _header_row_index(ws, contains='محصول')
        headers = _header_labels(ws, header_row)
        self.assertIn('محصول',      headers)
        self.assertIn('تامین‌کننده', headers)
        self.assertIn('تعداد',      headers)
        self.assertIn('مبلغ کل (تومان)', headers)

    def test_total_price_column_correct(self):
        r = self._get()
        ws = self._items_sheet(r.content)
        header_row = _header_row_index(ws, contains='محصول')
        headers = _header_labels(ws, header_row)
        col_idx = headers.index('مبلغ کل (تومان)') + 1
        total   = ws.cell(header_row + 1, col_idx).value
        # qty=3, unit_price=100 → total=300
        self.assertAlmostEqual(float(total), 300.0)

    def test_vendor_filter_applied(self):
        # Create second vendor with its own purchase
        v2   = _vendor('تامین‌کننده ب')
        p2   = _purchase(v2, status=PurchaseStatus.CONFIRMED)
        prod2 = _product('کالا دو')
        _purchase_item(p2, prod2)

        r = self._get(vendor=self.vendor.id)
        ws = self._items_sheet(r.content)
        header_row  = _header_row_index(ws, contains='محصول')
        vendor_col  = _header_labels(ws, header_row).index('تامین‌کننده') + 1
        data_rows   = _column_values(ws, header_row, vendor_col)
        self.assertTrue(data_rows)
        for name in data_rows:
            self.assertEqual(name, 'تامین‌کننده آ')

    def test_status_filter_applied(self):
        pend = _purchase(self.vendor, status=PurchaseStatus.PENDING)
        _purchase_item(pend, self.prod)

        r = self._get(status='CONFIRMED')
        ws = self._items_sheet(r.content)
        header_row  = _header_row_index(ws, contains='محصول')
        status_col  = _header_labels(ws, header_row).index('وضعیت خرید') + 1
        statuses    = _column_values(ws, header_row, status_col)
        self.assertTrue(statuses)
        for s in statuses:
            self.assertIn('تأیید', s)


# ---------------------------------------------------------------------------
# Product Cost Report export
# ---------------------------------------------------------------------------

class ProductCostReportExcelTest(APITestCase):

    def setUp(self):
        self.user = _admin()
        self.client.force_authenticate(self.user)
        vendor = _vendor()
        prod   = _product(purchase_price=Decimal('200'))
        purchase = _purchase(vendor, status=PurchaseStatus.CONFIRMED)
        _purchase_item(purchase, prod, qty=Decimal('5'), unit_price=Decimal('200'))

    def _get(self, **params):
        return self.client.get(COST_REPORT_URL, {'export': 'excel', **params})

    def test_returns_xlsx_for_finance_user(self):
        r = self._get(group_by='product')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], XLSX_CT)

    def test_product_group_by_headers(self):
        r = self._get(group_by='product')
        ws = _open_workbook(r.content).active
        headers = [ws.cell(row, col).value for row in range(1, ws.max_row + 1)
                   for col in range(1, ws.max_column + 1) if ws.cell(row, col).value]
        self.assertIn('محصول',       headers)
        self.assertIn('مجموع هزینه', headers)

    def test_vendor_group_by_headers(self):
        r = self._get(group_by='vendor')
        ws = _open_workbook(r.content).active
        headers = [ws.cell(row, col).value for row in range(1, ws.max_row + 1)
                   for col in range(1, ws.max_column + 1) if ws.cell(row, col).value]
        self.assertIn('تامین‌کننده', headers)

    def test_unauthorized_inventory_user_gets_403(self):
        u = User.objects.create_user(username='inv_usr', password='x')
        g, _ = Group.objects.get_or_create(name='inventory_user')
        u.groups.add(g)
        self.client.force_authenticate(u)
        r = self._get()
        self.assertEqual(r.status_code, status.HTTP_403_FORBIDDEN)

    def test_meta_rows_in_workbook(self):
        r = self._get(start_date='2025-01-01', end_date='2025-12-31')
        ws = _open_workbook(r.content).active
        # Meta rows are written before the header; check the first row is a label
        self.assertIsNotNone(ws.cell(1, 1).value)
